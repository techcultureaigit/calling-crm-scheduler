import re
import io
import csv
import struct
import httpx
import urllib.parse
from typing import List
import openpyxl
from app.core.config import settings
from app.core.logger import logger




async def upload_file_to_cloudinary(file_bytes: bytes, public_id: str, folder: str = "crm-call-recordings", resource_type: str = "video") -> str:
    """
    Uploads raw file bytes to Cloudinary.
    Returns the secure URL of the uploaded file.
    """
    import cloudinary
    import cloudinary.uploader

    cloud_name = getattr(settings, "CLOUDINARY_CLOUD_NAME", "not_set")
    api_key = getattr(settings, "CLOUDINARY_API_KEY", "not_set")
    api_secret = getattr(settings, "CLOUDINARY_API_SECRET", "not_set")

    if not all(v and v != "not_set" for v in [cloud_name, api_key, api_secret]):
        raise RuntimeError("Cloudinary credentials are not configured.")

    cloudinary.config(
        cloud_name=cloud_name,
        api_key=api_key,
        api_secret=api_secret,
        secure=True
    )

    file_obj = io.BytesIO(file_bytes)
    file_obj.name = f"{public_id}.mp3"

    logger.info(f"Uploading file to Cloudinary (folder: {folder}, public_id: {public_id}, size: {len(file_bytes)} bytes)...")

    result = cloudinary.uploader.upload(
        file_obj,
        public_id=public_id,
        resource_type=resource_type,
        folder=folder,
        overwrite=True
    )

    secure_url = result.get("secure_url")
    logger.info(f"Cloudinary upload successful. URL: {secure_url}")
    return secure_url


async def download_file_from_cloudinary(url: str) -> bytes:
    """
    Downloads raw contact file from Cloudinary.
    First tries direct download, falls back to signing the request if credentials are set.
    """
    logger.info(f"Downloading file from Cloudinary: {url}")
    
    # 1. Try direct HTTP GET first
    async with httpx.AsyncClient() as client:
        try:
            logger.info("Trying direct download...")
            response = await client.get(url, timeout=30.0)
            if response.status_code == 200:
                logger.info("Direct download successful.")
                return response.content
            logger.warning(f"Direct download returned status code {response.status_code}")
        except Exception as e:
            logger.warning(f"Direct download failed: {e}")

    # 2. Fall back to credentials-signed URL if config is available
    api_key = getattr(settings, "CLOUDINARY_API_KEY", "not_set")
    api_secret = getattr(settings, "CLOUDINARY_API_SECRET", "not_set")
    
    if api_key != "not_set" and api_secret != "not_set" and api_key and api_secret:
        try:
            logger.info("Generating signed Cloudinary URL for download...")
            import cloudinary
            import cloudinary.utils
            
            # Parse URL parts
            parsed_url = urllib.parse.urlparse(url)
            path_parts = [p for p in parsed_url.path.split("/") if p]
            
            if len(path_parts) >= 4:
                cloud_name = path_parts[0]
                resource_type = path_parts[1]  # 'raw', 'image', 'video'
                delivery_type = path_parts[2]  # 'upload', 'private', 'authenticated'
                
                version = None
                public_id_start_idx = 3
                if path_parts[3].startswith('v') and path_parts[3][1:].isdigit():
                    version = path_parts[3]
                    public_id_start_idx = 4
                    
                public_id = "/".join(path_parts[public_id_start_idx:])
                
                cloud_name_env = getattr(settings, "CLOUDINARY_CLOUD_NAME", cloud_name) or cloud_name
                
                cloudinary.config(
                    cloud_name=cloud_name_env,
                    api_key=api_key,
                    api_secret=api_secret,
                    secure=True
                )
                
                # Generate signed URL
                signed_url, options = cloudinary.utils.cloudinary_url(
                    public_id,
                    resource_type=resource_type,
                    type=delivery_type,
                    version=version[1:] if version else None,
                    sign_url=True
                )
                logger.info(f"Signed URL generated: {signed_url}")
                
                async with httpx.AsyncClient() as client:
                    response = await client.get(signed_url, timeout=30.0)
                    if response.status_code == 200:
                        logger.info("Signed download successful.")
                        return response.content
                    logger.error(f"Signed download returned status code {response.status_code}")
        except Exception as e:
            logger.error(f"Failed to download using signed URL: {e}")
            
    raise RuntimeError(f"Could not download contact file from Cloudinary URL: {url}")

def parse_contact_file(file_bytes: bytes, filename: str) -> List[str]:
    """
    Parses a contact list file (Excel or CSV format).
    Extracts the first column, starting from row 2 (skipping heading).
    """
    logger.info(f"Parsing contact file: {filename} ({len(file_bytes)} bytes)")
    
    # Check filename extension
    name_lower = filename.lower()
    numbers = []
    
    if name_lower.endswith(".xlsx") or name_lower.endswith(".xls"):
        # Excel parsing
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active
        # Iterate through rows starting from row 2
        for row in sheet.iter_rows(min_row=2, min_col=1, max_col=1, values_only=True):
            val = row[0]
            if val is not None:
                num_str = str(val).strip()
                # Clean up numeric strings
                if num_str.endswith(".0"):
                    num_str = num_str[:-2]
                # Remove common non-digits like space, +, -, ( )
                num_cleaned = re.sub(r"\D", "", num_str)
                if num_cleaned:
                    numbers.append(num_cleaned)
    else:
        # Fallback to CSV
        text_content = file_bytes.decode("utf-8", errors="ignore")
        reader = csv.reader(io.StringIO(text_content))
        first_row = True
        for row in reader:
            if first_row:
                first_row = False
                continue
            if row:
                val = row[0].strip()
                # Remove common non-digits
                num_cleaned = re.sub(r"\D", "", val)
                if num_cleaned:
                    numbers.append(num_cleaned)
                    
    logger.info(f"Successfully extracted {len(numbers)} contacts from file.")
    return numbers

import tempfile
import asyncio

async def stream_contact_file(url: str, filename: str):
    """
    Generator that yields contact numbers one by one without loading the entire
    file into memory. Supports CSV and Excel (.xlsx).
    """
    logger.info(f"Streaming contact file: {url}")
    name_lower = filename.lower()
    is_excel = name_lower.endswith(".xlsx") or name_lower.endswith(".xls")

    if is_excel:
        # Excel files cannot be easily streamed purely over HTTP, so we download to a temporary file
        # and use openpyxl's read_only=True mode.
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_file:
            tmp_path = tmp_file.name
        
        try:
            logger.info(f"Downloading Excel to temp file {tmp_path} for streaming...")
            async with httpx.AsyncClient() as client:
                async with client.stream("GET", url, timeout=60.0) as response:
                    if response.status_code != 200:
                        raise RuntimeError(f"Failed to download file from {url}. Status: {response.status_code}")
                    with open(tmp_path, "wb") as f:
                        async for chunk in response.aiter_bytes(chunk_size=8192):
                            f.write(chunk)
            
            logger.info("Opening Excel for streaming...")
            wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
            sheet = wb.active
            
            first_row = True
            for row in sheet.iter_rows(min_col=1, max_col=1, values_only=True):
                if first_row:
                    first_row = False
                    continue
                val = row[0]
                if val is not None:
                    num_str = str(val).strip()
                    if num_str.endswith(".0"):
                        num_str = num_str[:-2]
                    num_cleaned = re.sub(r"\D", "", num_str)
                    if num_cleaned:
                        yield num_cleaned
            
            wb.close()
        finally:
            import os
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
    else:
        # CSV processing using httpx stream
        logger.info("Streaming CSV file directly...")
        async with httpx.AsyncClient() as client:
            async with client.stream("GET", url, timeout=60.0) as response:
                if response.status_code != 200:
                    raise RuntimeError(f"Failed to download file from {url}. Status: {response.status_code}")
                
                first_row = True
                async for line in response.aiter_lines():
                    if not line.strip():
                        continue
                    if first_row:
                        first_row = False
                        continue
                        
                    # Basic CSV parsing (first column)
                    parts = line.split(",")
                    if parts:
                        val = parts[0].strip(' "\'')
                        num_cleaned = re.sub(r"\D", "", val)
                        if num_cleaned:
                            yield num_cleaned
